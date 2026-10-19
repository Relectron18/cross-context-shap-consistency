"""Response letter for the fourth-round concerns (Concerns4) -> paper/Response_to_Reviewer_v12.docx (earlier labels: v8-v10)"""
import sys; sys.path.insert(0, ".")
from pathlib import Path
import docx
from docx.shared import Pt
import resnums as N, json
A7 = json.load(open(N.RES / "p7_audit.json"))
from resnums import f, sg, ref, P1, P5, TK2

d = docx.Document()
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
d.add_heading("Response to the reviewer", 0)
d.add_paragraph("Manuscript: Cross-Context Consistency of SHAP-Based Feature Importance for Application Classification "
                "Across VPN and Non-VPN Traffic (revised version v12)")
d.add_paragraph("We thank the reviewer for another careful report. This revision adds a calibrated grouping design, a second grouping proxy based on the row order of the original ARFF files, block-summed "
                "SHAP shares with a conditional permutation check, a list-origin by target decomposition of the feature-list "
                "transfer, a family-share table and spreads for the class-mix swap. All revision history has been moved out of "
                "the manuscript and is collected at the end of this letter. Section and table numbers refer to the revised "
                "manuscript, which now has 15 tables and 8 figures.")

def C(title, quote, resp):
    d.add_heading(title, level=2)
    p = d.add_paragraph(); r = p.add_run("Reviewer: "); r.bold = True; p.add_run(quote).italic = True
    for i, t in enumerate(resp):
        p = d.add_paragraph()
        if i == 0: p.add_run("Response: ").bold = True
        p.add_run(t)

fpc = {c: N.flows_per_capture(c) for c in ("novpn", "vpn")}
lo = lambda c: min(fpc[c].values()); hi = lambda c: max(fpc[c].values())
cal = lambda m, k, s="mean": ref(m, "calibrated", k, s)
tk = TK2.groupby(["model", "k"])[["match", "vpn_list_gain"]].mean()

d.add_heading("Remaining concerns", level=1)
C("1. The session proxy is uncalibrated, and the conclusion depends on the granularity",
  "A cheap estimate would be the number of source captures versus the number of flows... At the coarsest setting, half-model "
  "macro F1 is only 0.53–0.61 and RF Δ has SD 0.111... use \"cannot be attributed to tunnelling\" throughout.",
  [f"We followed the suggestion (Section IV-D, new paragraph \"Calibrating the grouping\"). The processed files carry no capture "
   f"identifier, but the released captures can be counted: the capture-file prefix list in a public, independent implementation of Deep Packet ([7], [56] in the manuscript) names "
   f"109 Non-VPN and 30 VPN capture files for CHAT, FT, MAIL, STREAMING and VOIP (plus one VPN P2P capture). Dividing each "
   f"class's flows in the 15 s files by its number of captures gives {lo('novpn'):.0f}–{hi('novpn'):.0f} flows per capture "
   f"without a VPN and {lo('vpn'):.0f}–{hi('vpn'):.0f} with one, about 80% of which falls in a training partition. Non-VPN "
   f"captures are therefore closest to the ~35-flow setting and VPN captures to the coarsest setting or beyond. The paper states "
   f"this and also lists why the estimate is rough (third-party prefix list, a single P2P capture, captures are not sessions, and "
   f"a 15 s timeout splits long sessions).",
   f"We also added a calibrated design in which each class gets as many clusters as it has capture files (Table VII, Fig. 7). "
   f"Its mean gap is {sg(cal('XGB','rho_gap'))} for XGBoost (SD {f(cal('XGB','rho_gap','sd'))}) and {sg(cal('RF','rho_gap'))} for "
   f"Random Forest (SD {f(cal('RF','rho_gap','sd'))}), with half-model macro F1 of {f(cal('XGB','f1_nv'),2)}/{f(cal('XGB','f1_vpn'),2)} "
   f"and {f(cal('RF','f1_nv'),2)}/{f(cal('RF','f1_vpn'),2)}.",
   "We agree that this is weak evidence: at the coarse and calibrated cluster settings the half-models are weak and the gap "
   "varies widely between repeats. Because clusters are built from the same features the models use, we also suspected that at "
   "coarse settings they separate regions of feature space rather than captures, so we added a second proxy that does not use "
   "the features.",
   f"Row-order blocks (Section IV-D, new paragraph \"Row-order blocks\"; Table VII; Fig. 7). In the original Scenario A2 ARFF "
   f"files, flows are sorted by class and then kept in the order the flow meter wrote them. Every deduplicated flow in our data "
   f"matches exactly one ARFF row, so its position is known. Over the deduplicated flows, the nearest of the 40 rows "
   f"within 20 positions lies at {f(min(v['win20_dedup'] for v in A7['locality_dedup'].values()),2)}–"
   f"{f(max(v['win20_dedup'] for v in A7['locality_dedup'].values()),2)} of the random distance, which suggests that many "
   f"flows of one capture sit together. The manuscript also states where this evidence is weak: for VPN FT, MAIL, P2P and VOIP, "
   f"adjacent rows are almost as different as random pairs, row position correlates with some features, and BROWSING and P2P "
   f"cannot be split by capture at all. We therefore formed halves from contiguous blocks of rows within each class: about 10, 35 "
   f"and 140 flows per block, as many blocks as capture files, and the first versus the second half of each class. Under this "
   f"proxy the half-models stay accurate, and the gap does not vanish: {sg(ref('XGB','orderedcal','rho_gap'))} for XGBoost and "
   f"{sg(ref('RF','orderedcal','rho_gap'))} for Random Forest with blocks calibrated to the capture counts, and "
   f"{sg(ref('XGB','orderedhalf','rho_gap'))} and {sg(ref('RF','orderedhalf','rho_gap'))} with the first-versus-second-half "
   f"split. The Spearman gap is at or above zero in all 11 repeats, but the manuscript now says that this is weak evidence: the "
   f"first-versus-second-half repeats reuse nearly the same halves, and the RBO and Top-5 gaps are negative in some of them.",
   f"The manuscript also reports that the near-zero calibrated-cluster mean for Random Forest averages opposite gaps "
   f"({sg(A7['onesided_RF_calibrated']['nv'])} on the Non-VPN side, {sg(A7['onesided_RF_calibrated']['vpn'])} on the VPN side), and it "
   f"states what the designs do not test: the permutation null covers only random halves, the outer 80/20 split is always "
   f"random, and within-context pairs are explained on the same test records while cross-context pairs are not.",
   "This changes the conclusion of the paper, and the abstract, Sections VI-C, VI-F, VII, X-A, XI and XII now report both "
   "proxies. The two disagree, and neither is a capture split: clusters may overcorrect, and blocks may undercorrect where a "
   "capture crosses a block boundary. Read together, they show that part of the random-split gap reflects related flows and that "
   "some excess probably remains. Its size cannot be pinned down without capture identifiers, and because the two contexts are "
   "different sets of captures it cannot be attributed to tunnelling. Splitting by capture file after regenerating flows from the "
   "raw PCAPs remains the way to settle it, and Section XI says so.",
   f"The overcorrected sentences are gone. \"They should not be read as an effect of the VPN\" and similar wording are replaced by "
   f"\"cannot be attributed to tunnelling\" throughout, and Section VI-C notes that the Random Forest gap is still positive "
   f"in all 11 repeats at about 35 flows per cluster ({f(ref('RF','grouped35','rho_gap'))}, minimum {f(ref('RF','grouped35','rho_gap','min'))})."])

b = lambda m, c: P5[f"block_{m}_{c}"]["block"]; bh = lambda m, c: P5[f"block_half_{m}_{c}"]["mean"]
C("2. The min-IAT credit-redistribution conclusion holds for XGBoost but not clearly for RF",
  "Please report block-summed SHAP shares in Table IX and temper the sentence for RF. Also, permutation importance is not a "
  "neutral arbiter with correlated features (ref. [34]).",
  [f"Table IX now has a column with the summed SHAP share of the three minima, and Section VI-E opens with it. The reviewer's "
   f"reading is confirmed: for XGBoost the block keeps its share ({b('XGB','novpn'):.1f}% without a VPN and {b('XGB','vpn'):.1f}% "
   f"with one), while for Random Forest it grows from {b('RF','novpn'):.1f}% to {b('RF','vpn'):.1f}% (half-split means "
   f"{bh('RF','novpn'):.1f}% and {bh('RF','vpn'):.1f}%). The summary sentence now says that for XGBoost the visible shift is "
   f"largely redistribution of credit among correlated timing features, whereas for Random Forest the block itself gains about "
   f"six percentage points under VPN, so part of that shift is a change in what the model relies on.",
   "On permutation importance, Section VI-E now states, citing [34], that marginal permutation importance evaluates the model "
   "off the data distribution when features are correlated and is not a neutral arbiter. We added a conditional permutation "
   "check in the spirit of Strobl et al. [36]: each minimum IAT is permuted within quintile cells of its two most correlated "
   "features. Its ranks for the three minima (2–4 in every model and context) are reported in Table IX next to the marginal "
   "ranks. Both permutation measures move the minima less than SHAP does, and the text presents this as a second view rather "
   "than as a verdict, since conditional permutation depends on the conditioning scheme. It also reports that over the whole "
   "ranking permutation importance agrees less across contexts than SHAP does, so the smaller shifts apply to the minimum "
   "IATs only."])

C("3. The feature-list transfer should be read as a 2×2 (list × target)",
  "So the \"cost of transfer\" at the VPN target mostly reflects that the Non-VPN list is a worse list overall... Report this "
  "interaction term, and soften \"a measurable cost\".",
  [f"Agreed. Section VIII-B and the new Table XIII decompose the transfer into list origin and target. For each model, k and "
   f"selection seed, the context-match term is the sum of the two own-minus-transferred contrasts, and the VPN-list advantage is "
   f"the mean over both targets of the VPN list's macro F1 minus the Non-VPN list's. For Random Forest the VPN-derived list is "
   f"better on both targets ({sg(tk.loc[('RF',10),'vpn_list_gain'])} at k = 10 and {sg(tk.loc[('RF',5),'vpn_list_gain'])} at "
   f"k = 5, averaged over targets and seeds), and the context-match term is small: {sg(tk.loc[('RF',10),'match'])} (Random Forest) "
   f"and {sg(tk.loc[('XGB',10),'match'])} (XGBoost) at k = 10, {sg(tk.loc[('RF',5),'match'])} and {sg(tk.loc[('XGB',5),'match'])} "
   f"at k = 5, and within 0.003 of zero at k = 15.",
   "\"A measurable cost\" has been removed. The text now says that most of the apparent per-target cost for the VPN Random Forest "
   "reflects a weaker Non-VPN list rather than a mismatch between list and context, and that the effect of matching a list to its "
   "context is small for lists of ten or more features. The summary table (Table XV), Section X-D and the conclusion use the same "
   "framing. The manuscript also notes the limited power of this test: the five selection seeds yield only one to three "
   "distinct lists per model, context and k, and each list is refitted once."])

C("4. The XGBoost family-level gap does not shrink under grouping",
  "It needs a sentence of discussion, and the abstract's \"gap shrinks substantially\" should be scoped to feature level.",
  [f"Section VII now has a paragraph on this (after Table XI). For Random Forest the summed-family gap falls from "
   f"{f(ref('RF','random','fam_gap'))} to {sg(ref('RF','grouped10','fam_gap'))} with clusters of about ten flows. For XGBoost it "
   f"does not shrink at that setting ({f(ref('XGB','random','fam_gap'))} versus {f(ref('XGB','grouped10','fam_gap'))}, and "
   f"similarly for the correlation clusters) and falls only at about 35 flows per cluster ({f(ref('XGB','grouped35','fam_gap'))}) "
   f"or under the calibrated grouping ({sg(ref('XGB','calibrated','fam_gap'))}). The paragraph says that part of the XGBoost "
   f"difference survives fine grouping at the family level although it shrinks at the feature level, so the redundancy reading "
   f"applies less clearly to XGBoost families.",
   f"The abstract now states the shrinkage only for the Spearman gap between feature rankings, and Section XII calls it the feature-level gap. "
   f"The family-level result is named in the summary of Section X-C. Section VII also reports that under the row-order blocks the "
   f"summed-family gaps stay positive for both models.",])

C("5. The model-transfer comparison is loose",
  "These protocols differ... and class mix and sessions are confounded in the transfer. Hedge it.",
  ["Section VIII-A no longer says that the contexts differ \"at least as much\" as distant clusters. It now says that the transfer "
   "scores fall in the same range as the cluster-grouped cross-validation but that the comparison is only indicative, names the "
   "protocol differences (50 clusters and 80% of one context's training data versus a full model of the other context), and "
   "states that the change of context also changes class mix and capture sessions, which this test cannot separate. The practical "
   "conclusion (a model must be retrained) is kept, because it does not depend on that comparison."])

d.add_heading("Minor and presentational points", level=1)
rf800 = P5["rf800_halfmodels"]
minor = [
 ("Revision history in the manuscript.", "Removed. The manuscript no longer refers to earlier versions in Sections III-D, IV-D, "
  "IV-G, VI-C, VI-D or IX (or anywhere else); a text search for \"previous version\", \"earlier version\" and \"earlier reading\" "
  "returns nothing. The notes that were there are collected below."),
 ("Abstract length.", "Shortened from about 300 to about 225 words. It keeps the headline contrast (random halves, the two "
  "grouping proxies, permutation null, model transfer and list transfer) and leaves the remaining numbers to "
  "the summary table, which is now Table XV."),
 ("Class-mix swap spreads and range.", "Table VIII now gives mean (SD) over the 3–5 draws for every entry, plus a column for the "
  "agreement between natural-mix and swapped-mix models of the same draw. The text now quotes the two ranges separately: "
  "0.924–0.977 for natural versus swapped models (last column) and 0.936–0.986 for agreement with the context's primary ranking, "
  "which match the table."),
 ("RF explained on 800 records.", f"Confirmed for half-models. We refitted eight Random Forest half-models (two repeats, both "
  f"contexts, both halves) and explained each on its full test partition and on the 800-record sample: the rankings agree at "
  f"ρ ≥ {f(min(r['rho'] for r in rf800),3)}, with the same five leading features in all {len(rf800)} cases. Section IV-D states this."),
 ("Bottom-10 range.", "The text now gives the means over seeds, 0.685–0.735, as in the table."),
 ("Family shares.", "A compact table of family shares with ranks for both models and contexts is added (Table X), and Section VII "
  "quotes its values."),
]
for a, t in minor:
    p = d.add_paragraph(style="List Bullet"); r = p.add_run(a + " "); r.italic = True; p.add_run(t)

d.add_heading("Revision notes moved out of the manuscript", level=1)
d.add_paragraph("These statements were previously in the manuscript and are recorded here instead.")
notes = [
 "Logistic Regression was used as a linear baseline in earlier versions. It was dropped because it was a weak baseline on these features; Section III-D now says only that no linear baseline is used.",
 f"All results were recomputed in a single environment (Python {P1['env']['python']}, scikit-learn {P1['env']['sklearn']}, XGBoost {P1['env']['xgboost']}, SHAP {P1['env']['shap']}). Random Forest results did not change; XGBoost single fits moved by at most 0.006 macro F1. In an earlier version, two tables reported values for the same XGBoost experiment that differed by 0.011, which this rerun removed.",
 "The training-bootstrap reference and its tables and figures were removed. Percentile intervals are now used everywhere.",
 "In the feature-list transfer, k = 15 was the originally planned cutoff. k = 5, 10 and 15 are now reported with equal weight.",
 "Redrawing the 80/20 partition in every repeat roughly doubled the SDs of the half-split gaps compared with the version that used one fixed partition.",
 "An earlier version drew a stronger conclusion from the sentinel comparison than the data support. It was withdrawn: Section IX now compares the encoding effect with random-half retraining variation.",
 "An earlier version read the cross-context ρ of about 0.85 for class-matched 1,330-flow models as support for a class-mix explanation. With a same-size within-context reference (about 0.95), that reading was withdrawn, and Section VI-D explains that small training sets agree more with everything.",
]
for t in notes: d.add_paragraph(t, style="List Bullet")

d.add_heading("Note on the manuscript file", level=1)
d.add_paragraph("The submission system flagged the previous file for an unsupported language and for characters from other "
                "alphabets. We checked the file character by character. It contained no hidden or substituted letters, but it "
                "did contain mathematical symbols typed as plain text: Greek letters such as ρ, τ, κ and Δ, Unicode minus signs "
                "(U+2212), non-breaking spaces and arrows. The style sheet also declared Chinese and Hindi for the East Asian and complex-script language slots. "
                "In this version every mathematical symbol is set as a Word equation object, minus signs, non-breaking spaces and "
                "arrows are replaced with their ASCII equivalents, and the document language is English (US) throughout. Outside "
                "equations, the running text now contains only ASCII characters, typographic quotes and dashes, and the accented "
                "letters of author names in the reference list. The wording of the manuscript is unchanged by this step.")

out = Path(__file__).resolve().parents[2] / "paper" / "Response_to_Reviewer_v12.docx"
d.save(str(out)); print("ok", out)
