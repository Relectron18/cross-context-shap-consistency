"""v20: repair the team's v19 'we'-removal edits (typos, two garbled sentences, one dropped sentence, en-IN tags)."""
import docx, re, sys
from docx.oxml.ns import qn
SRC, DST = sys.argv[1], sys.argv[2]
d = docx.Document(SRC)
P = d.paragraphs

def rep(p, old, new):
    runs = p.runs; full = ''.join(r.text for r in runs)
    assert full.count(old) == 1, (old, full.count(old))
    s = full.index(old); e = s + len(old); pos = 0; first = True
    for r in runs:
        a, b = pos, pos + len(r.text); pos = b
        if b <= s or a >= e: continue
        t = r.text
        if first:
            r.text = t[:max(s - a, 0)] + new + (t[e - a:] if e < b else ''); first = False
        else:
            r.text = t[e - a:] if e < b else ''

def find(start):
    m = [p for p in P if p.text.startswith(start)]; assert len(m) == 1, start; return m[0]

E = [
 ("With correlated features, attribution", [("None of this is done.", "This study does neither."), ("member feature.  This", "member feature. This")]),
 ("SHAP has been used before", [("narrower .", "narrower.")]),
 ("The input is the time-related", [(
   "This split was created specifically for the present study and not supplied with the dataset, referred to hereafter as the primary split, and the two parts it comprises named the primary partitions. Every processed record was also checked against a corresponding row in the original ARFF files.",
   "This split was created for the present study and is not supplied with the dataset; it is referred to hereafter as the primary split, and its two parts as the primary partitions. Every processed record was also verified to map back to a row of the original ARFF files.")]),
 ("BROWSING is the lone exception", [("all of them ;", "all of them;")]),
 ("Averaging over output classes", [(
   "so those in frequent classes get more representativity than they should have. Importance is therefore also measured as equal weighting over the true classes. Neither approach, however, captures the class-specific",
   "so frequent classes contribute more records. Importance is therefore also reported with the true classes weighted equally. Neither version, however, captures class-specific")]),
 ("Calibrating the grouping.", [("context’s", "context's")]),
 ("Row-order blocks.", [("each class .", "each class.")]),
 ("For five grouped designs", [("on its own .", "on its own.")]),
 ("XGBoost half-models are explained", [("replications.The", "replications. The")]),
 ("Table I sorts the 23 features", [("familiesBecause", "families. Because")]),
 ("Two features of this grouping", [("; Families are", "; families are"),
   ("are therefore also used.", "are therefore also used. Family-level agreement is computed on the primary fits and in every half-split repeat.")]),
 ("In the feature-list transfer", [("15 run for both models and", "15 are run for both models, and")]),
 ("Effect sizes and their spread", [("not reported For", "not reported. For")]),
 ("Permutation importance. Permutation", [("also reported .", "also reported.")]),
 ("How large is the encoding effect?", [("rankings ;", "rankings;")]),
 ("Benchmark age, reuse", [("technologies .", "technologies.")]),
]
for start, pairs in E:
    p = find(start)
    for o, n in pairs: rep(p, o, n)

# four empty spacer paragraphs before TABLE XIV -> page break before the caption
i = [k for k, p in enumerate(P) if p.text == 'TABLE XIV'][0]
for k in range(i - 1, i - 5, -1):
    assert P[k].text == ''; P[k]._p.getparent().remove(P[k]._p)
d.paragraphs[[k for k, p in enumerate(d.paragraphs) if p.text == 'TABLE XIV'][0]].paragraph_format.page_break_before = True

# drop en-IN language tags added in v19 (v18 had none)
n = 0
for el in d.element.body.iter(qn('w:lang')):
    if el.get(qn('w:val')) == 'en-IN':
        el.getparent().remove(el); n += 1
print('lang tags removed', n)
d.save(DST)
