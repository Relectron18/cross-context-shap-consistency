import docx, copy, re
d = docx.Document('paper/SHAP_VPN_revised_v17.docx')
P = {}
for p in d.paragraphs:
    m = re.match(r'\[(\d+)\]', p.text)
    if m: P[int(m.group(1))] = p
assert len(P) == 59
DOI = {8:'10.1145/3485447.3512217', 9:'10.1109/EuroSP48549.2020.00018', 10:'10.1109/ACCESS.2020.2988359',
 11:'10.1109/ACCESS.2022.3204171', 12:'10.1109/TNSM.2021.3098157', 15:'10.48550/arXiv.1806.08049',
 16:'10.1609/aaai.v33i01.33013681', 17:'10.48550/arXiv.2202.01602', 19:'10.48550/arXiv.2006.16234',
 21:'10.1109/INFCOMW.2019.8845315', 22:'10.1145/3460120.3484758', 23:'10.1145/3548606.3560609',
 24:'10.5220/0006105602530262', 27:'10.1145/3545574', 28:'10.1109/OJCOMS.2022.3188750',
 29:'10.1016/j.bdr.2022.100359', 31:'10.1515/9781400881970-018', 34:'10.1007/s11222-021-10057-z',
 37:'10.48550/arXiv.2106.12228', 38:'10.1007/s10618-022-00840-5', 40:'10.1080/01605682.2020.1865846',
 43:'10.2307/1412159', 44:'10.1093/biomet/30.1-2.81', 46:'10.1201/9780429246593', 47:'10.1214/09-AOS699',
 57:'10.1073/pnas.102102699', 58:'10.1145/2382577.2382579', 59:'10.1016/j.patter.2023.100804'}
# rebuild [55] and [56] using [54]'s run formatting (label, plain, italic)
tpl = [r._r.rPr for r in P[54].runs]
def rebuild(p, parts):
    for r in list(p.runs): r._r.getparent().remove(r._r)
    for text, k in parts:
        r = p.add_run(text)
        if tpl[k] is not None: r._r.insert(0, copy.deepcopy(tpl[k]))
rebuild(P[55], [('[55]',0), ('\tM. G. Kendall, “The treatment of ties in ranking problems,” ',1), ('Biometrika',2),
                (', vol. 33, no. 3, pp. 239-251, 1945, doi: 10.1093/biomet/33.3.239.',1)])
rebuild(P[56], [('[56]',0), ('\tmunhouiani, “Deep-Packet: PyTorch implementation of Deep Packet,” GitHub repository. Accessed: Oct. 15, 2026. [Online]. Available: https://github.com/munhouiani/Deep-Packet',1)])
for n, doi in DOI.items():
    p = P[n]; assert 'doi:' not in p.text, n
    last = p.runs[-1]; assert last.text.endswith('.'), (n, last.text)
    last.text = last.text[:-1] + ', doi: ' + doi + '.'
d.save('paper/SHAP_VPN_revised_v18.docx')
d2 = docx.Document('paper/SHAP_VPN_revised_v18.docx')
refs = [p.text for p in d2.paragraphs if re.match(r'\[\d+\]\t', p.text)]
print(len(refs), sum('doi:' in r for r in refs))
for r in refs:
    if 'doi:' not in r: print('NO DOI:', r[:90])
